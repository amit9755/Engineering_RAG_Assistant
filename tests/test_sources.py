"""Phase 2 document lifecycle tests using isolated registry and vector databases."""

from io import BytesIO
from unittest.mock import Mock, patch

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from langchain_core.embeddings import Embeddings

from src.sources.documents import DocumentSourceService
from src.sources.models import BitbucketSourceConfig, SourceStatus
from src.sources.registry import SourceRegistry


class TestEmbeddings(Embeddings):
    """Small deterministic vectors keep lifecycle tests offline."""
    def embed_documents(self, texts):
        return [[float(len(text)), 1.0, 0.0] for text in texts]

    def embed_query(self, text):
        return self.embed_documents([text])[0]


@pytest.fixture
def service(tmp_path):
    from langchain_community.vectorstores import Chroma
    with patch("src.ingestion.embedder.EmbeddingModel._load_model"):
        from src.retrieval.vector_store import VectorStore
    vectors = object.__new__(VectorStore)
    vectors._store = Chroma(
        collection_name="phase2-tests", embedding_function=TestEmbeddings(),
        persist_directory=str(tmp_path / "vectors"),
    )
    return DocumentSourceService(
        registry=SourceRegistry(tmp_path / "registry" / "sources.db"),
        storage_dir=tmp_path / "originals", vector_store=vectors, refresh=Mock(),
    )


@pytest.fixture
def client(service, monkeypatch):
    from src.api.main import app
    from src.api.routes.sources import documents
    import src.sources.documents as service_module
    monkeypatch.setattr(documents, "document_service", service)
    monkeypatch.setattr(service_module, "document_service", service)
    with TestClient(app) as client:
        yield client


def test_upload_reindex_and_delete_preserve_other_sources(service):
    first = service.upload(b"First document. " * 100, "first.txt")
    second = service.upload(b"Second document", "second.txt")
    collection = service.vectors._store._collection
    count = collection.count()
    assert first.status == SourceStatus.READY
    assert service._path(first).read_bytes() == b"First document. " * 100
    chunks = collection.get(where={"source_id": first.id}, include=["metadatas"])
    assert len(chunks["ids"]) == first.chunk_count
    assert len({meta["chunk_id"] for meta in chunks["metadatas"]}) == first.chunk_count
    assert all(meta["source_file"] == "first.txt" and meta["source_type"] == "document"
               for meta in chunks["metadatas"])
    assert service.reindex(first.id).chunk_count == first.chunk_count
    assert collection.count() == count  # Replacement never accumulates duplicates.
    service.delete(first.id)
    assert not service._path(first).exists()
    assert service.registry.get_source(first.id) is None
    assert collection.get(where={"source_id": first.id})["ids"] == []
    assert collection.count() == second.chunk_count
    assert service._refresh.call_count == 4


def test_same_filename_has_independent_sources_and_safe_paths(service):
    one = service.upload(b"First", "../../resume.txt")
    two = service.upload(b"Second", "resume.txt")
    assert one.id != two.id
    assert one.filename == two.filename == "resume.txt"
    assert service._path(one).parent == service.storage_dir
    assert service._path(one).read_bytes() == b"First"
    assert service._path(two).read_bytes() == b"Second"


def test_reindex_failure_preserves_old_vectors_and_retry_clears_error(service):
    source = service.upload(b"Original text", "notes.txt")
    collection = service.vectors._store._collection
    before = collection.get(where={"source_id": source.id})
    real_add = collection.add

    def partial_failure(ids, documents, metadatas, embeddings):
        real_add(ids=ids[:1], documents=documents[:1], metadatas=metadatas[:1], embeddings=embeddings[:1])
        raise RuntimeError("Simulated write failure")

    with patch.object(type(collection), "add", side_effect=partial_failure):
        with pytest.raises(RuntimeError):
            service.reindex(source.id)
    assert collection.get(where={"source_id": source.id})["ids"] == before["ids"]
    assert service.get(source.id).status == SourceStatus.ERROR
    recovered = service.reindex(source.id)
    assert recovered.status == SourceStatus.READY
    assert recovered.error_message is None


def test_delete_failure_retains_registry_and_original(service):
    source = service.upload(b"Keep this document", "keep.txt")
    with patch.object(service.vectors, "delete_by_source_id", side_effect=RuntimeError("unavailable")):
        with pytest.raises(RuntimeError):
            service.delete(source.id)
    assert service.get(source.id).id == source.id
    assert service._path(source).is_file()


def test_original_survives_service_restart(service):
    source = service.upload(b"Persistent document", "notes.md")
    restarted = DocumentSourceService(
        registry=SourceRegistry(service.registry._db_path), storage_dir=service.storage_dir,
        vector_store=service.vectors, refresh=Mock(),
    )
    assert restarted.reindex(source.id).chunk_count == source.chunk_count


def test_document_api_lifecycle(client, service):
    response = client.post("/api/v1/sources/documents", files={"file": ("notes.txt", b"Test notes")})
    assert response.status_code == 201
    source = response.json()
    assert source["status"] == "ready"
    url = f'/api/v1/sources/documents/{source["id"]}'
    assert client.get(url).json()["filename"] == "notes.txt"
    assert len(client.get("/api/v1/sources/documents").json()) == 1
    assert client.post(url + "/reindex").json()["chunks_added"] == source["chunk_count"]
    assert client.delete(url).status_code == 200
    assert client.get(url).status_code == 404
    assert service.vectors._store._collection.count() == 0


def test_docx_legacy_upload_also_registers(client, service):
    document = WordDocument()
    document.add_paragraph("Word document through the sidebar upload.")
    buffer = BytesIO()
    document.save(buffer)
    response = client.post("/api/v1/ingest/file", files={"file": ("notes.docx", buffer.getvalue())})
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert service.get(data["source_id"]).filename == "notes.docx"


@pytest.mark.parametrize("endpoint", ["/api/v1/sources/documents", "/api/v1/ingest/file"])
def test_upload_validation(client, monkeypatch, endpoint):
    from src.api.routes import ingest
    assert client.post(endpoint, files={"file": ("bad.exe", b"binary")}).status_code == 400
    assert client.post(endpoint, files={"file": ("empty.txt", b"")}).status_code == 422
    monkeypatch.setattr(ingest, "MAX_FILE_BYTES", 10)
    assert client.post(endpoint, files={"file": ("large.txt", b"x" * 11)}).status_code == 413


def test_wrong_source_type_cannot_be_deleted_or_reindexed(client, service):
    service.registry.create_bitbucket_source(
        BitbucketSourceConfig(source_id="bb-test", workspace="team", repository="repo",
                              credential_id="test-reference"), name="Repository",
    )
    url = "/api/v1/sources/documents/bb-test"
    assert client.delete(url).status_code == 404
    assert client.post(url + "/reindex").status_code == 404
    assert service.registry.get_source("bb-test") is not None


def test_missing_original_and_invalid_docx_are_reported(client, service):
    source = service.upload(b"Notes", "notes.txt")
    service._path(source).unlink()
    assert client.post(f"/api/v1/sources/documents/{source.id}/reindex").status_code == 409
    response = client.post("/api/v1/sources/documents", files={"file": ("bad.docx", b"not a docx")})
    assert response.status_code == 422
    assert "No readable text" in response.json()["detail"]
    failed = [item for item in service.registry.list_document_sources() if item.filename == "bad.docx"][0]
    assert failed.status == SourceStatus.ERROR


def test_deleting_last_document_clears_keyword_index(service, monkeypatch):
    with patch("src.retrieval.hybrid_retriever.SemanticReRanker._load_model"):
        from src.retrieval.hybrid_retriever import hybrid_retriever
    from src.ingestion.ingestion_pipeline import IngestionPipeline
    import src.retrieval.vector_store as vector_module

    monkeypatch.setattr(vector_module, "vector_store", service.vectors)
    # Restore the singleton's original index after the test.
    from src.retrieval.hybrid_retriever import BM25Retriever
    monkeypatch.setattr(hybrid_retriever, "bm25", BM25Retriever())
    service._refresh = lambda: IngestionPipeline()._refresh_bm25_index(strict=True)
    source = service.upload(b"Bluetooth supports wireless communication.", "notes.txt")
    assert hybrid_retriever.bm25.search("Bluetooth")
    service.delete(source.id)
    assert hybrid_retriever.bm25.search("Bluetooth") == []


def test_source_chunks_are_not_collapsed_by_fusion(service):
    with patch("src.retrieval.hybrid_retriever.SemanticReRanker._load_model"):
        from src.retrieval.hybrid_retriever import hybrid_retriever
    from langchain_core.documents import Document
    chunks = [Document(page_content=f"Section {i}", metadata={"document_id": "doc-one", "chunk_id": str(i)})
              for i in range(3)]
    merged = hybrid_retriever._reciprocal_rank_fusion(chunks, [(chunks[0], 1.0)])
    assert len(merged) == 3
