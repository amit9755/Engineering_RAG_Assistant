from docx import Document as WordDocument

from src.ingestion.document_loader import DocumentLoader


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "sample.docx"
    document = WordDocument()
    document.add_paragraph("Bluetooth engineering notes.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Protocol"
    table.cell(0, 1).text = "Bluetooth LE"
    document.save(path)

    chunks = DocumentLoader().load_and_chunk(str(path))

    text = "\n".join(chunk.page_content for chunk in chunks)
    assert "Bluetooth engineering notes." in text
    assert "Protocol" in text
    assert "Bluetooth LE" in text
    assert all(chunk.metadata["source_file"] == "sample.docx" for chunk in chunks)
